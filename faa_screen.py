"""Screen structures on the current FAA OE/AAA Pre-Screening Tool."""

from __future__ import annotations

import base64
import re
import shutil
import tempfile
import time
from pathlib import Path

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from math import ceil

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

# ---------------------------------------------------------------------------
# Black & Veatch workbook and printout folder (work computer).
# ---------------------------------------------------------------------------
WORKBOOK_PATH = r"C:\Users\And137460\OneDrive - Black & Veatch\PG&E Sacramento & LA OHTL - PGE Projects and Files\Projects\Sobrante\Working\30% Design\74066820 - Sobrante-Grizzly-Claremont #2 FAA Screening.xlsm"
PDF_FOLDER = r"C:\Users\And137460\OneDrive - Black & Veatch\PG&E Sacramento & LA OHTL - PGE Projects and Files\Projects\Sobrante\Working\30% Design\#2 FAA"
# ---------------------------------------------------------------------------

DATUM = "NAD83"
SHEET = "FAA Criteria Tool"
DATA_SHEET = "Data Sheet"
MAX_FILE_ROWS = 250
PRESCREEN_URL = "https://oeaaa.faa.gov/oeaaa/oe3a/main/#/noticePrescreen"
STRUCTURE_TYPE_ID = "TRANSMISSION_LINE$T_L_TOWER"
DOWNLOADS_DIR = ""
PROFILE_DIR = Path(__file__).resolve().parent / "FAA_edge_profile"


@dataclass
class SheetMap:
    number: str = "B"
    pls: str = "C"
    lon_deg: str = "D"
    lat_deg: str = "E"
    elevation: str = "F"
    height: str = "G"
    lon_dms: str = "H"
    lat_dms: str = "I"
    traverseway: str = "J"
    on_airport: str = "K"
    result: str = "L"
    timestamp: str = "M"
    hashcol: str = "N"


SHEET_MAP = SheetMap()


def detect_sheet_map(ws) -> SheetMap:
    headers: dict[str, str] = {}
    for col in range(1, 20):
        raw = ws.cell(1, col).value
        if raw is None:
            continue
        headers[" ".join(str(raw).lower().split())] = get_column_letter(col)

    def pick(*needles: str, default: str) -> str:
        for text, letter in headers.items():
            if all(needle in text for needle in needles):
                return letter
        return default

    if not any("structure number" in text or "latitude" in text for text in headers):
        return SheetMap(
            number="A",
            pls="B",
            lon_deg="C",
            lat_deg="D",
            elevation="E",
            height="F",
            lon_dms="G",
            lat_dms="H",
            traverseway="I",
            on_airport="J",
            result="K",
            timestamp="L",
            hashcol="M",
        )
    return SheetMap(
        number=pick("structure number", default="B"),
        pls=pick("pls", default="C"),
        lon_deg=pick("longitude", "deg", default="D"),
        lat_deg=pick("latitude", "deg", default="E"),
        elevation=pick("elevation", default="F"),
        height=pick("structure height", default="G"),
        lon_dms=pick("longitude", "dms", default="H"),
        lat_dms=pick("latitude", "dms", default="I"),
        traverseway=pick("traverseway", default="J"),
        on_airport=pick("on airport", default="K"),
        result=pick("result", default="L"),
        timestamp=pick("time", default="M"),
        hashcol=get_column_letter(column_index_from_string(pick("time", default="M")) + 1),
    )
RESULT_RE = re.compile(
    r"(required to file|not required to file|does not exceed|exceed)",
    re.I,
)


@dataclass
class Structure:
    row: int
    number: str
    pls_file: str
    lon_deg: float | None
    lat_deg: float | None
    elevation_ft: float
    height_ft: float
    lon_dms: str
    lat_dms: str
    traverseway: str
    on_airport: str
    structure_type: str = "TRANSMISSION_LINE$T_L_TOWER"
    last_result: str = ""
    last_hash: str = ""
    used_elevation: int | None = None

    @property
    def elevation_whole(self) -> int:
        return int(round_half_up(self.elevation_ft))

    @property
    def height_whole(self) -> int:
        if float(self.height_ft).is_integer():
            return int(self.height_ft)
        return int(ceil(self.height_ft))


def round_half_up(value: float) -> int:
    return int(value + 0.5) if value >= 0 else int(value - 0.5)


def _as_float(value) -> float | None:
    if value is None or value == "":
        return None
    return float(value)


def parse_dms(text: str) -> tuple[str, str]:
    raw = str(text).strip()
    hemi = raw[-1].upper() if raw[-1:] in "NSEWnsew" else ""
    body = raw[:-1] if hemi else raw
    match = re.search(
        r"(-?\d+(?:\.\d+)?)\s*[dD°]\s*(\d+(?:\.\d+)?)\s*['′]\s*(\d+(?:\.\d+)?)",
        body,
    )
    if match:
        deg = int(float(match.group(1)))
        minutes = int(float(match.group(2)))
        seconds = round(float(match.group(3)), 2)
        return f"{abs(deg)}-{minutes:02d}-{seconds:05.2f}", hemi
    hyphen = re.match(
        r"(-?\d+)\s*-\s*(\d+)\s*-\s*(\d+(?:\.\d+)?)\s*([NSEW])?$",
        raw,
        re.I,
    )
    if hyphen:
        deg = int(hyphen.group(1))
        minutes = int(hyphen.group(2))
        seconds = float(hyphen.group(3))
        hemi = (hyphen.group(4) or hemi).upper()
        return f"{abs(deg)}-{minutes:02d}-{seconds:05.2f}", hemi
    raise ValueError(f"Unrecognized DMS value: {text!r}")


def load_structures(workbook_path: Path) -> list[Structure]:
    global SHEET_MAP
    workbook_path = Path(workbook_path)
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[SHEET]
    SHEET_MAP = detect_sheet_map(ws)
    cols = SHEET_MAP
    rows: list[Structure] = []
    for r in range(2, ws.max_row + 1):
        number = ws[f"{cols.number}{r}"].value
        if number is None or str(number).strip() == "":
            continue
        elev_val = ws[f"{cols.elevation}{r}"].value
        height_val = ws[f"{cols.height}{r}"].value
        if elev_val in (None, "") or height_val in (None, ""):
            continue
        rows.append(
            Structure(
                row=r,
                number=str(number).strip(),
                pls_file=str(ws[f"{cols.pls}{r}"].value or "").strip(),
                lon_deg=_as_float(ws[f"{cols.lon_deg}{r}"].value),
                lat_deg=_as_float(ws[f"{cols.lat_deg}{r}"].value),
                elevation_ft=float(elev_val),
                height_ft=float(height_val),
                lon_dms=str(ws[f"{cols.lon_dms}{r}"].value).strip(),
                lat_dms=str(ws[f"{cols.lat_dms}{r}"].value).strip(),
                traverseway=str(ws[f"{cols.traverseway}{r}"].value or "No Traverseway").strip(),
                on_airport=str(ws[f"{cols.on_airport}{r}"].value or "No").strip(),
                last_result=str(ws[f"{cols.result}{r}"].value or "").strip(),
                last_hash=str(ws[f"{cols.hashcol}{r}"].value or "").strip(),
            )
        )
    return rows


def _norm_dms(text: str) -> str:
    body, hemi = parse_dms(text)
    return f"{body}{hemi}".upper()


def content_key(structure: Structure) -> str:
    """Identity used to decide whether FAA must be queried again."""
    payload = "|".join(
        [
            _norm_dms(structure.lat_dms),
            _norm_dms(structure.lon_dms),
            str(structure.height_whole),
            str(structure.elevation_whole),
        ]
    )
    return sha256(payload.encode("utf-8")).hexdigest()[:16]


def legacy_fingerprint(structure: Structure, datum: str) -> str:
    payload = "|".join(
        [
            structure.number,
            str(structure.lat_dms),
            str(structure.lon_dms),
            str(structure.lat_deg),
            str(structure.lon_deg),
            str(structure.height_ft),
            str(structure.on_airport),
            str(structure.structure_type),
            datum,
            "USGS",
        ]
    )
    return sha256(payload.encode("utf-8")).hexdigest()[:16]


def input_fingerprint(structure: Structure, datum: str) -> str:
    return content_key(structure)


def write_result(
    workbook_path: Path,
    structure: Structure,
    result_text: str,
    requires_filing: bool,
) -> None:
    workbook_path = Path(workbook_path)
    wb = load_workbook(workbook_path, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    ws = wb[SHEET]
    cols = SHEET_MAP
    ws[f"{cols.result}{structure.row}"] = result_text
    ws[f"{cols.timestamp}{structure.row}"] = datetime.now()
    ws[f"{cols.hashcol}{structure.row}"] = content_key(structure)

    data = wb[DATA_SHEET]
    if requires_filing:
        lat_fmt, _ = parse_dms(structure.lat_dms)
        lon_fmt, _ = parse_dms(structure.lon_dms)
        next_row = 2
        while next_row <= MAX_FILE_ROWS + 1 and data[f"A{next_row}"].value:
            if str(data[f"A{next_row}"].value).strip() == structure.number:
                break
            next_row += 1
        if next_row > MAX_FILE_ROWS + 1:
            raise RuntimeError("Data Sheet already has 250 cases.")
        data[f"A{next_row}"] = structure.number
        data[f"B{next_row}"] = lat_fmt
        data[f"C{next_row}"] = lon_fmt
        data[f"D{next_row}"] = (
            structure.used_elevation
            if structure.used_elevation is not None
            else structure.elevation_whole
        )
        data[f"E{next_row}"] = structure.height_whole
    wb.save(workbook_path)


def safe_structure_name(number: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "-", str(number).strip())
    return cleaned.strip(" .-") or "structure"


def pdf_path(folder: Path, number: str, requires_filing: bool) -> Path:
    folder = Path(folder)
    suffix = "NEED TO FILE" if requires_filing else "FILE NOT REQUIRED"
    return folder / f"{safe_structure_name(number)}_FAA Screening_{suffix}.pdf"


def remove_stale_pdfs(folder: Path, number: str, keep: Path | None = None) -> None:
    folder = Path(folder)
    prefix = f"{safe_structure_name(number)}_FAA Screening_"
    for path in folder.glob(f"{prefix}*.pdf"):
        if keep is None or path.resolve() != keep.resolve():
            path.unlink(missing_ok=True)


def existing_named_pdf(folder: Path, number: str) -> Path | None:
    for filing in (True, False):
        path = pdf_path(folder, number, filing)
        if path.exists():
            return path
    return None


def read_pdf_fields(path: Path) -> tuple[str, str, int, int] | None:
    try:
        text = path.read_bytes().decode("latin-1", errors="ignore")
    except OSError:
        return None
    tokens = re.findall(r"\(([^\)]{1,80})\) Tj", text)
    try:
        lat_i = tokens.index("Latitude")
        lat, lon, height, elev = tokens[lat_i + 5 : lat_i + 9]
        return lat.strip(), lon.strip(), int(float(height)), int(float(elev))
    except (ValueError, IndexError):
        return None


def dms_equal(left: str, right: str) -> bool:
    try:
        return _norm_dms(left) == _norm_dms(right)
    except ValueError:
        return re.sub(r"\s+", "", left).upper() == re.sub(r"\s+", "", right).upper()


def pdf_matches_structure(path: Path, structure: Structure) -> bool:
    fields = read_pdf_fields(path)
    if not fields:
        return False
    lat, lon, height, _elev = fields
    return (
        dms_equal(lat, structure.lat_dms)
        and dms_equal(lon, structure.lon_dms)
        and height == structure.height_whole
    )


def pdf_is_reusable(path: Path, structure: Structure) -> bool:
    fields = read_pdf_fields(path)
    if fields is None:
        return True
    return pdf_matches_structure(path, structure)


def find_matching_pdf(
    structure: Structure, folder: Path, structures: list[Structure]
) -> Path | None:
    own = existing_named_pdf(folder, structure.number)
    if own and pdf_is_reusable(own, structure):
        return own
    key = content_key(structure)
    for other in structures:
        if other.number == structure.number:
            continue
        if content_key(other) != key:
            continue
        found = existing_named_pdf(folder, other.number)
        if found and pdf_is_reusable(found, structure):
            return found
    for path in folder.glob("*_FAA Screening_*.pdf"):
        if pdf_matches_structure(path, structure):
            return path
    return None


def find_prior_result(
    structure: Structure, structures: list[Structure]
) -> str:
    key = content_key(structure)
    if structure.last_result:
        return structure.last_result
    for other in structures:
        if other.number != structure.number and content_key(other) == key and other.last_result:
            return other.last_result
    return ""


def place_existing_pdf(
    src: Path, structure: Structure, folder: Path, live_numbers: set[str]
) -> Path:
    filing = "NEED TO FILE" in src.name
    dest = pdf_path(folder, structure.number, filing)
    if src.resolve() == dest.resolve():
        return dest
    src_stem = src.name.split("_FAA Screening_")[0]
    donor_still_used = any(safe_structure_name(num) == src_stem for num in live_numbers)
    if dest.exists() and dest.resolve() != src.resolve():
        dest.unlink()
    if donor_still_used:
        shutil.copy2(src, dest)
        print(f"  copied existing printout {src.name} -> {dest.name}")
    else:
        src.replace(dest)
        print(f"  renamed existing printout {src.name} -> {dest.name}")
    return dest


def refresh_hash_only(workbook_path: Path, structure: Structure) -> None:
    workbook_path = Path(workbook_path)
    if structure.last_hash == content_key(structure):
        return
    wb = load_workbook(workbook_path, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    ws = wb[SHEET]
    ws[f"{SHEET_MAP.hashcol}{structure.row}"] = content_key(structure)
    wb.save(workbook_path)


def decide_action(
    structure: Structure, structures: list[Structure], folder: Path
) -> tuple[str, Path | None]:
    """Return ('skip' | 'reuse' | 'screen', optional existing pdf)."""
    matching = find_matching_pdf(structure, folder, structures)
    own = existing_named_pdf(folder, structure.number)
    if own and matching is not None and own.resolve() == matching.resolve():
        return "skip", own
    if matching is not None:
        return "reuse", matching
    return "screen", None


def dismiss_overlays(page) -> None:
    for name in ("Stay Logged In", "Continue", "OK", "Close"):
        btn = page.get_by_role("button", name=name)
        try:
            if btn.count() and btn.first.is_visible():
                btn.first.click(timeout=1500)
                time.sleep(0.3)
        except Exception:
            pass


def hide_modals(page) -> None:
    page.evaluate(
        """() => {
            if (typeof $ !== 'undefined') {
                $('#verifyNoticePointModal').modal('hide');
                $('#confirmPointModal').modal('hide');
            }
        }"""
    )


def select_structure_type(page) -> None:
    page.locator("structure-type").click()
    page.wait_for_timeout(500)
    leaf = page.locator(f'span[data-id="{STRUCTURE_TYPE_ID}"]')
    if leaf.count() == 0:
        page.get_by_text("T-L Tower", exact=False).first.click()
        return
    leaf.first.scroll_into_view_if_needed()
    leaf.first.click()
    page.wait_for_timeout(400)


def apply_nad83_datum(page) -> None:
    page.evaluate(
        """(datum) => {
            const el = document.querySelector('notice-criteria');
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            const point = scope.data.points.find(p => p.include) || scope.data.points[0];
            point.datum = datum;
            if (scope.formData) scope.formData.datum = datum;
            if (scope.noticeData && scope.noticeData.verifyPointModalPoint) {
                scope.noticeData.verifyPointModalPoint.datum = datum;
            }
            scope.$apply();
        }""",
        DATUM,
    )


def read_accepted_elevation(page) -> tuple[int | None, str]:
    payload = page.evaluate(
        """() => {
            const el = document.querySelector('notice-criteria');
            if (!el || typeof angular === 'undefined') return null;
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            if (!scope) return null;
            const point = scope.data.points.find(p => p.include) || scope.data.points[0];
            if (!point) return null;
            return {
                elev: point.siteElevation,
                source: point.siteElevationSource || ''
            };
        }"""
    )
    if not payload:
        return None, ""
    elev = payload.get("elev")
    try:
        elev_int = int(round(float(elev))) if elev is not None and elev != "" else None
    except (TypeError, ValueError):
        elev_int = None
    return elev_int, str(payload.get("source") or "")


def fill_point(page, structure: Structure) -> None:
    lat_body, lat_dir = parse_dms(structure.lat_dms)
    lon_body, lon_dir = parse_dms(structure.lon_dms)
    if not lat_dir:
        lat_dir = "N"
    if not lon_dir:
        lon_dir = "W"

    page.evaluate(
        """([lat, latDir, lon, lonDir, height, onAirport, datum]) => {
            if (!document.getElementById('terrain-working-indicator-notice')) {
                const marker = document.createElement('div');
                marker.id = 'terrain-working-indicator-notice';
                marker.style.display = 'none';
                document.body.appendChild(marker);
            }
            const el = document.querySelector('notice-criteria');
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            const point = scope.data.points.find(p => p.include) || scope.data.points[0];
            point.formLat2 = lat;
            point.latDir = latDir;
            point.formLon2 = lon;
            point.lonDir = lonDir;
            point.datum = datum;
            point.structureHeight = Number(height);
            point.siteElevation = undefined;
            point.siteElevationSource = undefined;
            if (point.elevationDetails) {
                point.elevationDetails.comments = '';
            }
            point.state = 'COORD_DIRTY';
            scope.data.onAirport = !!onAirport;
            if (scope.formData) scope.formData.datum = datum;
            scope.$apply();
            return true;
        }""",
        [
            lat_body,
            lat_dir,
            lon_body,
            lon_dir,
            structure.height_whole,
            structure.on_airport.lower() == "yes",
            DATUM,
        ],
    )
    page.wait_for_timeout(800)
    validate = page.get_by_role("button", name="VALIDATE")
    validate.first.wait_for(state="visible", timeout=10000)
    validate.first.click()
    handle_point_modal(page, structure)
    apply_nad83_datum(page)
    if structure.lat_deg is not None and structure.lon_deg is not None:
        page.evaluate(
            """([lat, lon, datum]) => {
                const el = document.querySelector('notice-criteria');
                const scope = angular.element(el).isolateScope() || angular.element(el).scope();
                const point = scope.data.points.find(p => p.include) || scope.data.points[0];
                if (!point.lat) point.lat = lat;
                if (!point.lon) point.lon = lon;
                point.datum = datum;
                if (scope.formData) scope.formData.datum = datum;
                if (scope.isPointValid) scope.isPointValid(point);
                scope.$apply();
            }""",
            [structure.lat_deg, structure.lon_deg, DATUM],
        )
    page.wait_for_function(
        """() => {
            const el = document.querySelector('notice-criteria');
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            const point = scope.data.points.find(p => p.include) || scope.data.points[0];
            return point && point.siteElevation !== undefined && point.siteElevation !== null
                && point.siteElevation !== '' && String(point.siteElevationSource || '').toUpperCase() !== 'USER';
        }""",
        timeout=10000,
    )
    elev, source = read_accepted_elevation(page)
    if elev is None or source.upper() == "USER":
        raise RuntimeError(
            f"FAA did not keep the NAD83/USGS site elevation (elev={elev!r}, source={source!r})"
        )
    structure.used_elevation = elev
    page.wait_for_timeout(400)


def handle_point_modal(page, structure: Structure) -> None:
    modal = page.locator("#verifyNoticePointModal")
    confirm = page.locator("#confirmPointModal")
    try:
        page.wait_for_function(
            """() => {
                const isOpen = (id) => {
                    const el = document.querySelector(id);
                    if (!el) return false;
                    const style = window.getComputedStyle(el);
                    return style.display !== 'none' && style.visibility !== 'hidden'
                        && (el.classList.contains('show') || el.classList.contains('in'));
                };
                return isOpen('#verifyNoticePointModal') || isOpen('#confirmPointModal');
            }""",
            timeout=25000,
        )
    except PlaywrightTimeout:
        raise RuntimeError(
            "FAA point validation dialog did not open after VALIDATE. "
            "The NAD83/USGS site elevation was not available."
        )
    if confirm.count() and confirm.first.is_visible() and not (
        modal.count() and modal.first.is_visible()
    ):
        page.get_by_role("button", name="Ok").click()
        page.wait_for_timeout(400)
        return
    page.wait_for_function(
        """() => {
            const el = document.querySelector('notice-criteria');
            if (!el || typeof angular === 'undefined') return false;
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            const elev = scope && scope.noticeData && scope.noticeData.verifyPointModalPoint
                && scope.noticeData.verifyPointModalPoint.usgsSiteElev;
            return elev !== undefined && elev !== null && elev !== '';
        }""",
        timeout=20000,
    )
    keep = page.locator("#customRadioInline1")
    if keep.count():
        keep.check()
    page.evaluate(
        """(datum) => {
            const el = document.querySelector('notice-criteria');
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            if (scope.noticeData && scope.noticeData.verifyPointModalPoint) {
                scope.noticeData.verifyPointModalPoint.siteElevationSource = 'USGS';
                scope.noticeData.verifyPointModalPoint.datum = datum;
            }
            if (scope.formData) scope.formData.datum = datum;
            scope.$apply();
        }""",
        DATUM,
    )
    page.get_by_role("button", name="Accept Point").click()
    page.wait_for_timeout(800)
    apply_nad83_datum(page)


def read_scope_results(page):
    return page.evaluate(
        """() => {
            const el = document.querySelector('notice-criteria');
            if (!el || typeof angular === 'undefined') return null;
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            if (!scope) return null;
            return {
                exceed: scope.preScreenResults ? scope.preScreenResults.isExceedence : null,
                results: scope.results || [],
                header: scope.resultHeader || ''
            };
        }"""
    )


def compose_result_text(payload: dict) -> str:
    exceed = payload.get("exceed") or []
    messages = [str(m).strip() for m in (payload.get("results") or []) if str(m).strip()]
    header = str(payload.get("header") or "").strip()
    if any(bool(x) for x in exceed):
        lead = (
            header
            or "Based on the information you provided, you are required to file notice with the FAA."
        )
    elif isinstance(exceed, list) and len(exceed) > 0:
        lead = (
            header
            or "Based on the information you provided, you are not required to file notice with the FAA."
        )
    else:
        lead = header or ""
    if not lead and messages:
        lead = " ".join(messages)
    parts = [lead] + [m for m in messages if m not in lead]
    return "\n".join(p for p in parts if p)


def submit_and_read(page) -> str:
    invoked = page.evaluate(
        """() => {
            const el = document.querySelector('notice-criteria');
            const scope = angular.element(el).isolateScope() || angular.element(el).scope();
            if (!scope || !scope.noticeCriteria) return false;
            scope.noticeCriteria();
            scope.$applyAsync();
            return true;
        }"""
    )
    if not invoked:
        page.get_by_role("button", name="Submit").click()

    deadline = time.time() + 60
    payload = None
    while time.time() < deadline:
        payload = read_scope_results(page)
        exceed = payload.get("exceed") if payload else None
        if isinstance(exceed, list) and len(exceed) > 0:
            break
        body = page.locator("body").inner_text()
        if RESULT_RE.search(body) and "Pre-Screening Tool" in body:
            if "required to file" in body.lower() or "does not exceed" in body.lower():
                break
        page.wait_for_timeout(500)
    else:
        raise RuntimeError(f"FAA pre-screen did not return results: {payload!r}")

    text = compose_result_text(payload or {})
    if not RESULT_RE.search(text):
        visible = page.locator("body").inner_text()
        matches = [ln.strip() for ln in visible.splitlines() if RESULT_RE.search(ln)]
        text = "\n".join(matches) or text
    if not RESULT_RE.search(text):
        raise RuntimeError(f"No FAA result text. scope={payload!r}")
    return " ".join(text.split())


def requires_filing(result_text: str) -> bool:
    lower = result_text.lower()
    if "not required to file" in lower or "does not exceed" in lower:
        return False
    if "required to file" in lower or "exceed" in lower:
        return True
    return False


def _install_pdf_capture(page) -> None:
    page.evaluate(
        """() => {
            if (window.__faaPdfCaptureInstalled) return;
            window.__faaPdfCaptureInstalled = true;
            window.__faaPdfBase64 = null;
            const remember = (blob) => {
                if (!blob || window.__faaPdfBase64) return;
                const type = (blob.type || '').toLowerCase();
                if (type && type.indexOf('pdf') === -1 && type !== '') return;
                blob.arrayBuffer().then((buf) => {
                    const bytes = new Uint8Array(buf);
                    let bin = '';
                    const chunk = 0x8000;
                    for (let i = 0; i < bytes.length; i += chunk) {
                        bin += String.fromCharCode.apply(null, bytes.subarray(i, i + chunk));
                    }
                    const encoded = btoa(bin);
                    if (encoded.indexOf('JVBER') === 0) {
                        window.__faaPdfBase64 = encoded;
                    }
                });
            };
            const orig = URL.createObjectURL;
            URL.createObjectURL = function(obj) {
                try { remember(obj); } catch (err) {}
                return orig.apply(this, arguments);
            };
        }"""
    )


def _pdf_bytes_from_capture(page) -> bytes | None:
    try:
        page.wait_for_function("() => !!window.__faaPdfBase64", timeout=20000)
    except PlaywrightTimeout:
        return None
    encoded = page.evaluate("() => window.__faaPdfBase64")
    if not encoded:
        return None
    data = base64.b64decode(encoded)
    if data.startswith(b"%PDF") and len(data) >= 1000:
        return data
    return None


def save_result_pdf(page, dest: Path) -> None:
    dest = dest.resolve()
    dest.parent.mkdir(parents=True, exist_ok=True)
    hide_modals(page)
    print_btn = page.get_by_role("button", name="Print")
    print_btn.first.wait_for(state="visible", timeout=15000)
    _install_pdf_capture(page)

    download = None
    try:
        with page.expect_download(timeout=25000) as download_info:
            print_btn.first.click()
        download = download_info.value
    except PlaywrightTimeout:
        print("  no browser download event; reading the in-page Print PDF ...")
        if print_btn.first.is_visible():
            print_btn.first.click()

    captured = _pdf_bytes_from_capture(page)
    if captured:
        dest.write_bytes(captured)
        print(f"  filed {dest} ({dest.stat().st_size} bytes)")
        return

    if download is not None:
        try:
            if dest.exists():
                dest.unlink()
            download.save_as(str(dest))
        except Exception as exc:
            print(f"  save_as failed ({exc})")
            src = download.path()
            if src and Path(src).exists():
                shutil.copy2(src, dest)

    if (not dest.exists() or dest.stat().st_size < 1000) and DOWNLOADS_DIR:
        named = list(Path(DOWNLOADS_DIR).rglob("prescreen.pdf"))
        if named:
            shutil.copy2(named[0], dest)

    if not dest.exists() or dest.stat().st_size < 1000:
        raise RuntimeError(f"FAA Print PDF was not written: {dest}")
    print(f"  filed {dest} ({dest.stat().st_size} bytes)")


def open_browser(playwright, downloads_dir: Path):
    PROFILE_DIR.mkdir(parents=True, exist_ok=True)
    lock = PROFILE_DIR / "SingletonLock"
    if lock.exists():
        try:
            lock.unlink()
        except OSError:
            pass
    options = dict(
        headless=False,
        viewport={"width": 1400, "height": 900},
        accept_downloads=True,
        downloads_path=str(downloads_dir),
        chromium_sandbox=True,
    )
    try:
        return playwright.chromium.launch_persistent_context(
            str(PROFILE_DIR), channel="msedge", **options
        )
    except Exception as exc:
        print(f"  System Edge not available ({exc}); using Playwright Chromium.")
        return playwright.chromium.launch_persistent_context(str(PROFILE_DIR), **options)


def open_prescreen(page, reset_form: bool) -> None:
    dismiss_overlays(page)
    already_open = "noticePrescreen" in (page.url or "")
    if already_open and reset_form:
        clear = page.get_by_role("button", name="Clear")
        if clear.count() and clear.first.is_visible():
            clear.first.click()
            page.wait_for_timeout(400)
        else:
            already_open = False
    if not already_open:
        print("  opening Pre-Screening Tool ...")
        page.goto(PRESCREEN_URL, wait_until="domcontentloaded")
        dismiss_overlays(page)
    page.wait_for_selector("structure-type", timeout=30000)
    page.wait_for_function(
        """() => typeof angular !== 'undefined'
            && !!document.querySelector('notice-criteria')""",
        timeout=30000,
    )


def screen_one(page, structure: Structure, pdf_folder: Path, booted: bool) -> tuple[str, Path]:
    try:
        open_prescreen(page, reset_form=booted)

        print("  selecting structure type ...")
        select_structure_type(page)
        print("  filling NAD83 point and keeping FAA/USGS site elevation ...")
        fill_point(page, structure)
        print(f"  using site elevation {structure.used_elevation} ft ...")
        print("  submitting ...")
        result = submit_and_read(page)
        if not RESULT_RE.search(result):
            raise RuntimeError(f"No FAA result text for {structure.number}: {result[:300]!r}")
        apply_nad83_datum(page)
        hide_modals(page)
        filing = requires_filing(result)
        dest = pdf_path(pdf_folder, structure.number, filing)
        remove_stale_pdfs(pdf_folder, structure.number, keep=None)
        print("  printing official FAA PDF ...")
        save_result_pdf(page, dest)
        print(f"  saved {dest.name}")
        return result, dest
    except Exception as exc:
        print(f"  ERROR {structure.number}: {type(exc).__name__}: {exc}")
        try:
            if not page.is_closed():
                page.screenshot(
                    path=str(pdf_folder / f"{safe_structure_name(structure.number)}_FAILED.png"),
                    full_page=True,
                )
        except Exception as shot_exc:
            print(f"  Could not save failure screenshot: {shot_exc}")
        raise


def run(workbook: Path, pdf_folder: Path) -> list[tuple[Structure, str, bool]]:
    global DOWNLOADS_DIR
    workbook = Path(workbook)
    pdf_folder = Path(pdf_folder)
    pdf_folder.mkdir(parents=True, exist_ok=True)
    structures = load_structures(workbook)
    if not structures:
        raise SystemExit(f"No structure rows found in {workbook}")

    live_numbers = {item.number for item in structures}
    to_screen: list[Structure] = []
    outcomes: list[tuple[Structure, str, bool]] = []
    for structure in structures:
        action, existing = decide_action(structure, structures, pdf_folder)
        if action == "skip":
            print(f"Skipping {structure.number}: same lat/long/height/elevation already on file")
            refresh_hash_only(workbook, structure)
            continue
        if action == "reuse" and existing is not None:
            print(
                f"Reusing printout for {structure.number}: "
                "same lat/long/height/elevation, new structure number"
            )
            dest = place_existing_pdf(existing, structure, pdf_folder, live_numbers)
            result = find_prior_result(structure, structures)
            if not result:
                filing = "NEED TO FILE" in dest.name
                result = (
                    "Based on the information you provided, you are required to file notice with the FAA."
                    if filing
                    else "Based on the information you provided, you are not required to file notice with the FAA."
                )
            filing = requires_filing(result)
            write_result(workbook, structure, result, filing)
            outcomes.append((structure, result, filing))
            continue
        to_screen.append(structure)

    if not to_screen:
        print("Nothing new to screen.")
        return outcomes

    DOWNLOADS_DIR = tempfile.mkdtemp(prefix="faa_dl_")
    with sync_playwright() as playwright:
        context = open_browser(playwright, Path(DOWNLOADS_DIR))
        page = context.pages[0] if context.pages else context.new_page()
        page.on("dialog", lambda dialog: dialog.accept())
        try:
            booted = False
            for structure in to_screen:
                print(f"Screening {structure.number} ...")
                result, _dest = screen_one(page, structure, pdf_folder, booted)
                booted = True
                filing = requires_filing(result)
                write_result(workbook, structure, result, filing)
                outcomes.append((structure, result, filing))
                print(f"  {structure.number}: {result[:180]}")
                page.wait_for_timeout(2000)
        finally:
            context.close()
    return outcomes


def main() -> None:
    workbook = Path(WORKBOOK_PATH)
    pdf_folder = Path(PDF_FOLDER)
    if not workbook.exists():
        raise SystemExit(f"Workbook not found: {workbook}")
    print(f"Workbook: {workbook}")
    print(f"PDF folder: {pdf_folder}")
    outcomes = run(workbook, pdf_folder)
    print("\nDone.")
    for structure, result, filing in outcomes:
        status = "NEED TO FILE" if filing else "FILE NOT REQUIRED"
        print(f"{structure.number}\t{status}\t{result[:120]}")


if __name__ == "__main__":
    main()
